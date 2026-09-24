import json, re, sys, unicodedata, zipfile, tempfile, os
from datetime import datetime
from pathlib import Path
import openpyxl

def norm(v):
    s = unicodedata.normalize('NFD', str(v or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()

def num(v):
    if isinstance(v, (int, float)): return float(v or 0)
    s = str(v or '').strip().replace('.', '').replace(',', '.')
    try: return float(s) if s else 0
    except ValueError: return 0

canonical = {
    '19826':'Partage','20740':'Madureira','21044':'Alcantara','21469':'Juiz de Fora',
    '21470':'Benfica','21483':'Tres Rios','22552':'Raul Soares','22554':'Alem Paraiba',
    '22555':'Manhuacu','22588':'Leopoldina','23318':'Santos Dumont','23433':'Caratinga',
    '23441':'Carangola','24064':'Aimores'
}
roles = {'bronze','cobre','diamante','diamante gb','esmeralda gb','ouro','platina','prata','rubi','revendedor'}
def parse(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as first_error:
        # Algumas exportações do Excel trazem caracteres de controle inválidos no XML.
        # Remove apenas esses bytes proibidos e tenta ler uma cópia temporária.
        cleaned = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False).name
        try:
            with zipfile.ZipFile(path, 'r') as source, zipfile.ZipFile(cleaned, 'w', zipfile.ZIP_DEFLATED) as target:
                for info in source.infolist():
                    payload = source.read(info.filename)
                    if info.filename.endswith('.xml'):
                        payload = re.sub(rb'[\x00-\x08\x0b\x0c\x0e-\x1f]', b'', payload)
                    target.writestr(info, payload)
            wb = openpyxl.load_workbook(cleaned, read_only=True, data_only=True)
        except Exception:
            try: os.unlink(cleaned)
            except OSError: pass
            raise first_error
    if 'Pag' not in wb.sheetnames: raise RuntimeError('A planilha precisa conter a aba Pag.')
    ws = wb['Pag']; it = ws.iter_rows(values_only=True); header = next(it, None)
    if not header: raise RuntimeError('A aba Pag não contém linhas de dados.')
    idx = {norm(v): i for i, v in enumerate(header)}
    def col(*names):
        for name in names:
            if norm(name) in idx: return idx[norm(name)]
        return None
    c_source = col('CanalDistribuicao'); c_type = col('Tipo de Entrega')
    c_cancel = col('SituaçãoComercial'); c_reason = col('DetalheSituaçãoComercial','Detalhe Situacao Comercial')
    c_fiscal = col('SituaçãoFiscal','Situacao Fiscal'); c_role = col('Papel'); c_meio = col('MeioCaptacao')
    c_items = col('QtdeItens'); c_date = col('Data Captação')
    required = [c_source,c_type,c_cancel,c_role,c_items,c_date]
    if any(v is None for v in required): raise RuntimeError('A aba Pag não possui todas as colunas necessárias.')
    buckets = {}; daily = {}; dates = set()
    def new_bucket(code):
        return {'store':canonical[code],'storeCode':code,'total':0,'retirada':0,'entrega':0,'revendedor':0,'omni':0,'revendedorCategorias':{},'cancelamentoMotivos':{},'cancelamentoFiscal':{},'retiradaCancelados':0,'entregaCancelados':0,'itens':0}
    def apply_row(b, row, retirada, cancelled):
        if cancelled:
            b['retiradaCancelados' if retirada else 'entregaCancelados'] += 1
            d = norm(row[c_reason] if c_reason is not None else '')
            key = 'usuario' if 'pelo usuario' in d else 'prazoAnalisePagamento' if 'analise do pagamento excedido' in d else 'analisePagamento' if 'analise do pagamento' in d else 'antifraude' if 'antifraude' in d else 'inatividade' if 'inatividade' in d else 'estoque' if 'inconsistencia de estoque' in d else 'inconsistencia' if 'inconsistencia' in d else 'prazoPendencia' if 'pendencia excedido' in d else 'recusaExterna' if 'autorizacao externa' in d else 'outros'
            a=b['cancelamentoMotivos'].setdefault(key,[0,0]); a[1 if retirada else 0]+=1
            fk=norm(row[c_fiscal] if c_fiscal is not None else '')
            if fk:
                a=b['cancelamentoFiscal'].setdefault(fk,[0,0]); a[1 if retirada else 0]+=1
        else:
            b['total'] += 1; b['itens'] += num(row[c_items]); b['retirada' if retirada else 'entrega'] += 1
            role=norm(row[c_role]); meio=norm(row[c_meio] if c_meio is not None else '')
            if role == 'consumidor final':
                b['omni'] += 1
            else:
                b['revendedor'] += 1; cat='diamante' if role=='diamante gb' else role; b['revendedorCategorias'][cat]=b['revendedorCategorias'].get(cat,0)+1
    for row in it:
        source = str(row[c_source] or '')
        m = re.search(r'\b(\d{4,6})\b', source)
        code = m.group(1) if m else None
        if code not in canonical: continue
        day = str(row[c_date] or '')
        dm = re.search(r'(\d{2}/\d{2}/\d{4})', day)
        day_key = None
        if dm:
            day_key = f'{dm.group(1)[6:]}-{dm.group(1)[3:5]}-{dm.group(1)[:2]}'
            dates.add(day_key)
        typ = norm((row[c_type] if c_type is not None else '') or '') + ' ' + norm((row[col('Detalhe Entrega')] if col('Detalhe Entrega') is not None else '') or '')
        retirada = 'retirada' in typ or 'retirar na central de servicos' in typ or 'loja' in typ
        cancelled = 'cancelado' in norm(row[c_cancel])
        b = buckets.setdefault(code, new_bucket(code))
        apply_row(b, row, retirada, cancelled)
        if day_key:
            db = daily.setdefault(f'{code}|{day_key}', new_bucket(code))
            apply_row(db, row, retirada, cancelled)
    days=max(1,len(dates)); result=[]
    for b in buckets.values():
        b['itens']=int(b['itens']) if b['itens'].is_integer() else b['itens']
        b.update(pctEntrega=b['entrega']/b['total'] if b['total'] else 0,pctRetirada=b['retirada']/b['total'] if b['total'] else 0,mediaRetirada=b['retirada']/days,mediaEntrega=b['entrega']/days,mediaOmni=b['omni']/days,mediaItens=b['itens']/days)
        result.append(b)
    result.sort(key=lambda x:(-x['total'],x['store']))
    daily_result=[]
    for key, b in daily.items():
        day=key.split('|',1)[1]
        b['itens']=int(b['itens']) if b['itens'].is_integer() else b['itens']
        b.update(date=day,pctEntrega=b['entrega']/b['total'] if b['total'] else 0,pctRetirada=b['retirada']/b['total'] if b['total'] else 0)
        daily_result.append(b)
    st=Path(path).stat(); return {'source':{'path':str(Path(path).resolve()),'fileName':Path(path).name,'modifiedAt':datetime.fromtimestamp(st.st_mtime).isoformat(),'size':st.st_size},'period':{'start':min(dates) if dates else None,'end':max(dates) if dates else None,'days':len(dates)},'stores':result,'daily':daily_result}

if __name__ == '__main__': print(json.dumps(parse(sys.argv[1]), ensure_ascii=False, separators=(',',':')))
